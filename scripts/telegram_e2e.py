"""Private stdin-only fake transport harness. No web route and no real adapter."""

import asyncio
import json
import os
import sys
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps/api"))
import hub.models  # noqa: E402,F401
from hub.auth.models import Manager, RateBucket  # noqa: E402
from hub.auth.service import create_manager  # noqa: E402
from hub.core.config import Settings  # noqa: E402
from hub.telegram.delivery import deliver_one  # noqa: E402
from hub.telegram.transport import parse_update  # noqa: E402
from hub.telegram.updates import process_update  # noqa: E402
from hub.telegram.worker import Worker  # noqa: E402
from sqlalchemy import delete  # noqa: E402
from tests.fakes import BOT_ID, BOT_USERNAME, FakeTelegram  # noqa: E402


async def main():
    url = os.environ.get(
        "TEST_DATABASE_URL",
        "postgresql+asyncpg://hub:hub_test_only@127.0.0.1:5437/technician_hub_test",
    )
    parsed = make_url(url)
    if parsed.database != "technician_hub_test" or parsed.host not in {"127.0.0.1", "localhost"}:
        raise RuntimeError("Isolated local test database required")
    payload = json.load(sys.stdin)
    engine = create_async_engine(url, hide_parameters=True)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    provider = FakeTelegram()
    try:
        action = payload["action"]
        if action == "bootstrap":
            async with factory() as db:
                # Each serial browser test gets isolated rate-limit state. The CLI
                # guard above permits only the disposable loopback test database;
                # application login limits and production routes are unchanged.
                await db.execute(delete(RateBucket))
                await db.commit()
                await create_manager(db, payload["username"], payload["password"])
            settings = Settings(
                database_url=url,
                app_env="test",
                telegram_mode="fake",
                telegram_expected_bot_id=BOT_ID,
                telegram_expected_bot_username=BOT_USERNAME,
            )
            await Worker(settings, engine, provider).startup()
        elif action in {"work_report_issue", "expense_issue"}:
            from hub.telegram.types import TrustedEvent

            settings = Settings(
                database_url=url, app_env="test", allowed_origins=[payload["origin"]]
            )
            result = await process_update(
                factory,
                provider,
                TrustedEvent(
                    payload["update_id"],
                    "COMMAND",
                    user_id=payload["user_id"],
                    chat_id=payload["user_id"],
                    chat_type="private",
                    command="/expenses" if action == "expense_issue" else "/report",
                ),
                BOT_ID,
                settings=settings,
            )
            if result.outcome != (
                "EXPENSE_FORM" if action == "expense_issue" else "WORK_REPORT_FORM"
            ):
                raise RuntimeError("Test form was not issued")
            print(json.dumps({"url": result.reply.splitlines()[-1]}))
        elif action in {"work_report_cleanup", "expense_cleanup"}:
            from sqlalchemy import text

            async with factory() as db, db.begin():
                # Guarded disposable database only; product has no deletion bypass.
                await db.execute(
                    text(
                        "TRUNCATE technician_form_sessions, work_report_revisions, "
                        "work_reports, expense_revisions, technician_expenses"
                    )
                )
        elif action == "accounting_setup":
            from datetime import datetime, timedelta
            from uuid import UUID
            from zoneinfo import ZoneInfo

            from tests.accounting_data import seed

            today = datetime.now(ZoneInfo("America/Los_Angeles")).date()
            tid = UUID(payload["technician_id"])
            await seed(
                factory,
                today,
                technician_id=tid,
                reports=[
                    dict(day=0, amount="100.01", payment_method="CASH", google=3, facebook=2),
                    dict(
                        day=0,
                        amount="93.00",
                        payment_method="CREDIT_CARD",
                        groupon=1,
                        closed_by="CALL_CENTER",
                        maintenance=True,
                    ),
                ],
                expenses=[dict(day=0, amount="20.00"), dict(day=0, amount="20.00")],
            )
            await seed(
                factory,
                today - timedelta(days=today.weekday() + 7),
                technician_id=tid,
                reports=[dict(day=0, amount="55.00", payment_method="SUPER")],
                expenses=[],
            )
            print(json.dumps({"today": str(today)}))
        elif action == "accounting_correct":
            from uuid import UUID

            from sqlalchemy import text

            async with factory() as db, db.begin():
                await db.execute(text("SET LOCAL session_replication_role = replica"))
                identifier = await db.scalar(
                    text("""SELECT w.id FROM work_reports w
                    JOIN work_report_revisions r ON r.report_id=w.id AND
                    r.revision_number=w.current_revision_number
                    WHERE technician_id=:tid AND payment_method='CASH'"""),
                    {"tid": UUID(payload["technician_id"])},
                )
                await db.execute(
                    text("""INSERT INTO work_report_revisions SELECT
                    (jsonb_populate_record(NULL::work_report_revisions, to_jsonb(r) ||
                    jsonb_build_object('id',gen_random_uuid(),'revision_number',2,
                    'amount_closed',175.01))).* FROM work_report_revisions r
                    WHERE report_id=:id AND revision_number=1"""),
                    {"id": identifier},
                )
                await db.execute(
                    text("UPDATE work_reports SET current_revision_number=2 WHERE id=:id"),
                    {"id": identifier},
                )
        elif action == "mirror_google_setup":
            from datetime import UTC, datetime

            from hub.accounting_mirrors.provider import SHEETS_SCOPE
            from hub.core.secrets import SecretCipher
            from hub.google_calendar.models import CalendarConnection
            from hub.google_calendar.types import SCOPE
            from sqlalchemy import select

            key = "acpU14UT8v3AsP9osikC88Q27CgN1o0jLIZu4aPKwHE="
            async with factory() as db, db.begin():
                connection = await db.scalar(
                    select(CalendarConnection).where(CalendarConnection.is_current.is_(True))
                )
                if connection is None:
                    connection = CalendarConnection(
                        provider="GOOGLE",
                        account_key="stage9-e2e-primary",
                        account_label="Stage 9 fake Google",
                        is_current=True,
                        status="CONNECTED",
                        generation=1,
                        encrypted_refresh_token=SecretCipher(key).encrypt("fake-refresh-only"),
                        granted_scopes=[SCOPE, SHEETS_SCOPE],
                        connected_at=datetime.now(UTC),
                    )
                    db.add(connection)
                else:
                    connection.status = "CONNECTED"
                    connection.encrypted_refresh_token = SecretCipher(key).encrypt(
                        "fake-refresh-only"
                    )
                    connection.granted_scopes = [SCOPE, SHEETS_SCOPE]
            print(json.dumps({"ready": True}))
        elif action == "mirror_process":
            from hub.accounting_mirrors.models import AccountingMirrorTarget
            from hub.accounting_mirrors.worker import run_once
            from hub.google_calendar.fake import FakeCalendarProvider
            from hub.google_calendar.types import SCOPE, SHEETS_SCOPE, TokenGrant
            from pydantic import SecretStr
            from sqlalchemy import select

            fake = FakeCalendarProvider()
            fake.grant = TokenGrant(
                "fake-access-only",
                "fake-refresh-only",
                tuple(sorted((SCOPE, SHEETS_SCOPE))),
            )
            settings = Settings(
                database_url=url,
                app_env="test",
                google_mode="fake",
                google_calendar_credential_encryption_key=SecretStr(
                    "acpU14UT8v3AsP9osikC88Q27CgN1o0jLIZu4aPKwHE="
                ),
            )
            results = await run_once(factory, settings, fake, limit=20)
            async with factory() as db:
                targets = (await db.scalars(select(AccountingMirrorTarget))).all()
            inspected = []
            for target in targets:
                book = fake.spreadsheets.get(target.spreadsheet_id)
                if not book:
                    continue
                for sheet in book["sheets"]:
                    inspected.append(
                        {
                            "kind": target.kind,
                            "spreadsheet_id": target.spreadsheet_id,
                            "title": sheet["title"],
                            "values": sheet["values"],
                            "format_count": len(sheet["formats"]),
                            "rows": sheet["rows"],
                            "columns": sheet["columns"],
                        }
                    )
            print(json.dumps({"results": results, "calls": fake.sheets_calls, "sheets": inspected}))
        elif action == "mirror_cleanup":
            from hub.accounting_mirrors.models import AccountingMirrorTarget

            async with factory() as db, db.begin():
                await db.execute(delete(AccountingMirrorTarget))
        elif action == "xlsx_inspect":
            from datetime import datetime
            from decimal import Decimal

            from openpyxl import load_workbook

            root = Path(__file__).resolve().parents[1]
            workbook_path = Path(payload["path"]).resolve()
            if workbook_path.suffix.lower() != ".xlsx" or root not in workbook_path.parents:
                raise RuntimeError("Workbook inspection is limited to local test artifacts")
            book = load_workbook(workbook_path, data_only=False, keep_links=False)
            cells = [
                cell
                for worksheet in book.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            values = [cell.value for cell in cells]
            expected_total = Decimal(payload["expected_total"])
            expected_names = payload.get("expected_names", [payload["expected_name"]])
            expected_week_start = payload.get("expected_week_start")
            week_token = (
                datetime.fromisoformat(expected_week_start).strftime("%m/%d/%Y")
                if expected_week_start
                else None
            )
            first = book.active
            band_titles = [
                first.cell(1, column).value
                for column in range(1, first.max_column + 1, 13)
                if first.cell(1, column).value
            ]
            style_columns = [1 + index * 13 for index in range(len(expected_names))]
            print(
                json.dumps(
                    {
                        "sheet": book.active.title,
                        "name_found": payload["expected_name"] in values,
                        "expected_names_found": all(name in values for name in expected_names),
                        "band_titles": band_titles,
                        "expected_band_titles": [
                            name for name in band_titles if name in expected_names
                        ],
                        "week_found": week_token is None
                        or any(week_token in value for value in values if isinstance(value, str)),
                        "total_found": any(
                            isinstance(value, (int, float))
                            and Decimal(str(value)).quantize(Decimal("0.01")) == expected_total
                            for value in values
                        ),
                        "total_match_count": sum(
                            isinstance(value, (int, float))
                            and Decimal(str(value)).quantize(Decimal("0.01")) == expected_total
                            for value in values
                        ),
                        "expense_total_found": payload.get("expected_expense") is None
                        or Decimal(str(first["L38"].value)).quantize(Decimal("0.01"))
                        == Decimal(payload["expected_expense"]),
                        "cash_found": payload.get("expected_cash") is None
                        or Decimal(str(first["L44"].value)).quantize(Decimal("0.01"))
                        == Decimal(payload["expected_cash"]),
                        "google_reviews_found": payload.get("expected_google") is None
                        or first["L40"].value == payload["expected_google"],
                        "representative_styles": first["A2"].fill.fgColor.rgb == "FF7D98D3"
                        and first["A3"].fill.fgColor.rgb == "FFFFFF00"
                        and first["C4"].number_format == "$#,##0.00"
                        and first["L58"].number_format == "$#,##0.00",
                        "block_style_parity": all(
                            first.cell(2, column)._style == first["A2"]._style
                            and first.cell(3, column)._style == first["A3"]._style
                            for column in style_columns
                        ),
                        "formula_count": sum(cell.data_type == "f" for cell in cells),
                        "hyperlink_count": sum(cell.hyperlink is not None for cell in cells),
                        "external_links": len(book._external_links),
                        "size": workbook_path.stat().st_size,
                    }
                )
            )
        elif action == "expense_setup":
            from uuid import UUID

            from hub.integrations.models import TelegramBinding

            async with factory() as db, db.begin():
                db.add(
                    TelegramBinding(
                        technician_id=UUID(payload["technician_id"]),
                        bot_id=BOT_ID,
                        telegram_user_id=payload["user_id"],
                        private_status="CONNECTED",
                        private_availability="AVAILABLE",
                        private_generation=1,
                    )
                )
        elif action == "schedule_setup":
            from uuid import UUID

            from hub.calendars.models import Calendar
            from hub.integrations.models import TelegramBinding
            from sqlalchemy import select

            async with factory() as db, db.begin():
                calendar = await db.get(Calendar, UUID(payload["calendar_id"]))
                calendar.provider_calendar_id = "test-stable-schedule"
                binding = await db.get(TelegramBinding, UUID(payload["technician_id"]))
                if not binding:
                    binding = TelegramBinding(technician_id=UUID(payload["technician_id"]))
                    db.add(binding)
                binding.bot_id, binding.telegram_user_id = BOT_ID, payload["user_id"]
                binding.telegram_group_chat_id = payload["chat_id"]
                binding.private_status = binding.group_status = "CONNECTED"
                binding.private_availability = binding.group_availability = "AVAILABLE"
                binding.private_generation = binding.group_generation = (
                    binding.group_private_generation
                ) = 1
        elif action == "schedule_deliver":
            from hub.schedule_delivery.delivery import deliver_one as deliver_schedule
            from hub.schedule_delivery.models import ScheduleDispatch
            from hub.telegram.types import ProviderError, TrustedEvent
            from pydantic import SecretStr
            from sqlalchemy import select

            settings = Settings(
                database_url=url,
                app_env="test",
                telegram_mode="fake",
                telegram_expected_bot_id=BOT_ID,
                telegram_expected_bot_username=BOT_USERNAME,
                google_mode="fake",
                google_calendar_credential_encryption_key=SecretStr(
                    "acpU14UT8v3AsP9osikC88Q27CgN1o0jLIZu4aPKwHE="
                ),
                schedule_delivery_enabled=True,
                schedule_payload_encryption_key=SecretStr(
                    os.environ["SCHEDULE_PAYLOAD_ENCRYPTION_KEY"]
                ),
            )
            provider.group(payload["chat_id"], payload["user_id"], payload["user_id"])
            if payload.get("ambiguous"):
                provider.send_error = ProviderError("NETWORK_UNCERTAIN")
            await deliver_schedule(factory, engine, provider, settings)
            if payload.get("acknowledge"):
                async with factory() as db:
                    dispatch = await db.scalar(
                        select(ScheduleDispatch)
                        .where(ScheduleDispatch.status == "SENT")
                        .order_by(ScheduleDispatch.created_at.desc())
                        .limit(1)
                    )
                await process_update(
                    factory,
                    provider,
                    TrustedEvent(
                        100001,
                        "CALLBACK",
                        chat_id=dispatch.chat_id,
                        user_id=payload["user_id"],
                        message_id=dispatch.message_id,
                        payload=provider.schedule_sent[-1][2],
                        callback_query_id="test-query",
                    ),
                    BOT_ID,
                )
        elif action == "google_reset":
            from hub.calendars.models import Calendar, CalendarAssignment
            from hub.google_calendar.models import CalendarConnection, GoogleOAuthAttempt
            from sqlalchemy import select

            async with factory() as db, db.begin():
                google_ids = select(Calendar.id).where(Calendar.source == "GOOGLE")
                await db.execute(
                    delete(CalendarAssignment).where(CalendarAssignment.calendar_id.in_(google_ids))
                )
                await db.execute(delete(Calendar).where(Calendar.source == "GOOGLE"))
                await db.execute(delete(GoogleOAuthAttempt))
                await db.execute(delete(CalendarConnection))
        elif action == "cleanup":
            async with factory() as db, db.begin():
                await db.execute(delete(Manager).where(Manager.username == payload["username"]))
        elif action == "claim":
            link = urlparse(payload["link"])
            query = parse_qs(link.query)
            group = "startgroup" in query
            token = query["startgroup" if group else "start"][0]
            user = payload["user_id"]
            chat = payload["chat_id"] if group else user
            provider.group(chat, user, user)
            event = parse_update(
                {
                    "update_id": payload["update_id"],
                    "message": {
                        "message_id": 1,
                        "from": {"id": user, "is_bot": False, "first_name": "E2E Telegram"},
                        "chat": {
                            "id": chat,
                            "type": "supergroup" if group else "private",
                            "title": "E2E work group",
                        },
                        "text": f"/start@{BOT_USERNAME} {token}",
                    },
                },
                BOT_USERNAME,
            )
            result = await process_update(factory, provider, event, BOT_ID)
            if result.outcome != "CONNECTED":
                raise RuntimeError("Fake claim did not connect")
        elif action == "deliver":
            provider.group(payload["chat_id"], payload["user_id"], payload["user_id"])
            for _ in range(20):
                if not await deliver_one(factory, engine, provider, BOT_ID):
                    break
        else:
            raise RuntimeError("Unknown harness action")
    finally:
        await engine.dispose()


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except Exception:
        raise SystemExit("Isolated Telegram harness failed. No live provider was called.") from None
