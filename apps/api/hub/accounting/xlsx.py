import re
import unicodedata
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from hub.accounting.domain import AccountingTotals, WeeklyAccounting
from hub.accounting.xlsx_styles import (
    BLACK,
    BLOCK_COLUMN_WIDTHS,
    COLUMN_HEADER,
    COLUMN_HEADER_ROW_HEIGHT,
    DATA_ROW_HEIGHT,
    DAY_HEADER,
    DAY_HEADER_ROW_HEIGHT,
    EXPENSE_AMOUNT,
    EXPENSE_CELL,
    EXPENSE_HEADER,
    JOB_AMOUNT,
    JOB_CELL,
    SUMMARY_AMOUNT,
    SUMMARY_CELL,
    SUMMARY_HEADER,
    TECHNICIAN_TITLE,
    TITLE_ROW_HEIGHT,
    TOTAL_VALUE,
    WEEK_LABEL,
    WEEK_METADATA,
    WRAPPED_CELL,
    CellStyle,
    wrapped_row_height,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WEEKDAYS = ("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY")
JOB_HEADERS = (
    "#",
    "JOB / LOCATION",
    "AMOUNT",
    "PAYMENT METHOD",
    "WHO CLOSED PROJECT?",
    "GROUPON REVIEW",
    "GOOGLE REVIEW",
    "FACEBOOK REVIEW",
    "YEARLY MAINTENANCE PLAN PROVIDED",
)
PAYMENT_LABELS = {
    "CASH": "Cash",
    "ZELLE": "Zelle",
    "CHECK": "Check",
    "CREDIT_CARD": "Credit Card / Cash App",
    "VENMO": "Venmo",
    "SUPER": "SUPER",
    "ESTIMATE": "Estimate",
    "CANCEL": "Cancel",
}
REVIEW_LABELS = {"GOOGLE": "Google", "GROUPON": "Groupon", "FACEBOOK": "Facebook"}
CLOSER_LABELS = {"MYSELF": "Closed by Myself", "CALL_CENTER": "Closed by Call center"}
TECHNICIAN_BLOCK_COLUMNS = 12
TECHNICIANS_PER_BAND = 10
TECHNICIAN_COLUMN_STRIDE = 13
MIN_JOB_ROWS = 15
MIN_BLOCK_ROWS = 119
BAND_GAP_ROWS = 3
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_FILENAME_UNSAFE = re.compile(r"[^A-Za-z0-9_-]+")
_THIN = Side(style="thin", color=BLACK)
_GRID = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


@dataclass(frozen=True)
class XlsxReport:
    sequence: int
    title: str
    location: str
    amount: Decimal
    payment_method: str
    closed_by: str
    google_reviews: int
    groupon_reviews: int
    facebook_reviews: int
    maintenance: bool


@dataclass(frozen=True)
class XlsxExpense:
    business_date: date
    expense_type: str
    amount: Decimal
    note: str


@dataclass(frozen=True)
class XlsxDay:
    business_date: date
    reports: tuple[XlsxReport, ...]


@dataclass(frozen=True)
class XlsxTotals:
    gross_total: Decimal
    expense_total: Decimal
    report_count: int
    expense_count: int
    maintenance_count: int
    payments: tuple[tuple[str, Decimal], ...]
    reviews: tuple[tuple[str, int], ...]
    closed_by: tuple[tuple[str, int], ...]

    @classmethod
    def from_accounting(cls, totals: AccountingTotals) -> "XlsxTotals":
        return cls(
            gross_total=totals.gross_total,
            expense_total=totals.expense_total,
            report_count=totals.report_count,
            expense_count=totals.expense_count,
            maintenance_count=totals.maintenance_count,
            payments=tuple(totals.payments.items()),
            reviews=tuple(totals.reviews.items()),
            closed_by=tuple(totals.closed_by.items()),
        )


@dataclass(frozen=True)
class WeeklyAccountingXlsxModel:
    technician_id: UUID
    technician_name: str
    accounting_timezone: str | None
    week_start: date
    week_end: date
    days: tuple[XlsxDay, ...]
    expenses: tuple[XlsxExpense, ...]
    totals: XlsxTotals

    @classmethod
    def from_accounting(cls, accounting: WeeklyAccounting) -> "WeeklyAccountingXlsxModel":
        return cls(
            technician_id=accounting.technician_id,
            technician_name=accounting.technician_name,
            accounting_timezone=accounting.accounting_timezone,
            week_start=accounting.week_start,
            week_end=accounting.week_end,
            days=tuple(
                XlsxDay(
                    business_date=day.business_date,
                    reports=tuple(
                        XlsxReport(
                            sequence=report.sequence,
                            title=report.title,
                            location=report.location,
                            amount=report.amount,
                            payment_method=report.payment_method,
                            closed_by=report.closed_by,
                            google_reviews=report.google_reviews,
                            groupon_reviews=report.groupon_reviews,
                            facebook_reviews=report.facebook_reviews,
                            maintenance=report.maintenance,
                        )
                        for report in day.reports
                    ),
                )
                for day in accounting.days
            ),
            expenses=tuple(
                XlsxExpense(
                    business_date=expense.business_date,
                    expense_type=expense.expense_type,
                    amount=expense.amount,
                    note=expense.note,
                )
                for day in accounting.days
                for expense in day.expenses
            ),
            totals=XlsxTotals.from_accounting(accounting.totals),
        )


@dataclass(frozen=True)
class BlockLayout:
    day_rows: tuple[int, ...]
    summary_row: int
    row_count: int


def block_layout(model: WeeklyAccountingXlsxModel) -> BlockLayout:
    starts = []
    cursor = 0
    for day in model.days:
        starts.append(cursor)
        cursor += 2 + max(MIN_JOB_ROWS, len(day.reports))
    summary = max(35, len(model.expenses) + 1)
    return BlockLayout(tuple(starts), summary, max(MIN_BLOCK_ROWS, cursor, summary + 23))


def spreadsheet_text(value: object) -> str:
    return _INVALID_XML.sub("", str(value))[:32767]


def safe_filename_component(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return _FILENAME_UNSAFE.sub("_", normalized).strip("_-")[:80] or "technician"


def individual_filename(model: WeeklyAccountingXlsxModel) -> str:
    return (
        f"{safe_filename_component(model.technician_name)}_"
        f"{model.week_start.isoformat()}_{model.week_end.isoformat()}.xlsx"
    )


def all_tech_filename(week_start: date, week_end: date) -> str:
    return f"All_Tech_{week_start.isoformat()}_{week_end.isoformat()}.xlsx"


def _write(worksheet, row: int, column: int, value: object):
    cell = worksheet.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = spreadsheet_text(value)
        cell.data_type = "s"
        cell.hyperlink = None
    else:
        cell.value = value
    return cell


def _style(cell, style: CellStyle) -> None:
    cell.font = Font(
        name=style.font_name,
        size=style.font_size,
        bold=style.bold,
        color=f"FF{style.font_color}",
    )
    cell.fill = (
        PatternFill("solid", fgColor=f"FF{style.fill_color}") if style.fill_color else PatternFill()
    )
    cell.border = _GRID if style.bordered else Border()
    cell.alignment = Alignment(
        horizontal=style.horizontal,
        vertical=style.vertical,
        wrap_text=style.wrap_text,
    )
    if style.number_format:
        cell.number_format = style.number_format


def _height(worksheet, row: int, value: int) -> None:
    worksheet.row_dimensions[row].height = max(worksheet.row_dimensions[row].height or 0, value)


def _new_workbook(title: str, week_start: date):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    fixed = datetime.combine(week_start, time.min, tzinfo=UTC)
    workbook.properties.creator = "Technician Hub"
    workbook.properties.title = title
    workbook.properties.subject = "Monday-Sunday weekly accounting"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    return workbook, worksheet


def _configure(worksheet, *, all_tech: bool) -> None:
    worksheet.sheet_view.showGridLines = True
    worksheet.sheet_view.zoomScale = 55 if all_tech else 85
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15
    )


def _set_widths(worksheet, start_column: int) -> None:
    for offset, width in enumerate(BLOCK_COLUMN_WIDTHS):
        worksheet.column_dimensions[get_column_letter(start_column + offset)].width = width


def _title(worksheet, model: WeeklyAccountingXlsxModel, row: int, column: int) -> None:
    name = _write(worksheet, row, column, model.technician_name)
    _style(name, TECHNICIAN_TITLE)
    zone = _write(
        worksheet,
        row,
        column + 8,
        f"Timezone: {model.accounting_timezone or 'Not configured'}",
    )
    _style(zone, WEEK_METADATA)
    week = _write(worksheet, row, column + 10, "WEEK")
    _style(week, WEEK_LABEL)
    dates = _write(
        worksheet,
        row,
        column + 11,
        f"{model.week_start:%m/%d/%Y} - {model.week_end:%m/%d/%Y}",
    )
    _style(dates, WEEK_METADATA)
    _height(worksheet, row, TITLE_ROW_HEIGHT)


def _day_section(
    worksheet, model: WeeklyAccountingXlsxModel, day_index: int, row: int, column: int
) -> None:
    day = model.days[day_index]
    title = f"{WEEKDAYS[day_index]}  {day.business_date:%m/%d/%Y}"
    for offset in range(9):
        title_cell = _write(worksheet, row, column + offset, title if offset == 0 else "")
        _style(title_cell, DAY_HEADER)
        header = _write(worksheet, row + 1, column + offset, JOB_HEADERS[offset])
        _style(header, COLUMN_HEADER)
    _height(worksheet, row, DAY_HEADER_ROW_HEIGHT)
    _height(worksheet, row + 1, COLUMN_HEADER_ROW_HEIGHT)
    count = max(MIN_JOB_ROWS, len(day.reports))
    for index in range(count):
        current_row = row + 2 + index
        report = day.reports[index] if index < len(day.reports) else None
        values = (
            report.sequence if report else "",
            f"{report.title}\n{report.location}" if report else "",
            report.amount if report else "",
            PAYMENT_LABELS[report.payment_method] if report else "",
            CLOSER_LABELS[report.closed_by].removeprefix("Closed by ") if report else "",
            report.groupon_reviews if report else "",
            report.google_reviews if report else "",
            report.facebook_reviews if report else "",
            "Yes" if report and report.maintenance else "No" if report else "",
        )
        for offset, value in enumerate(values):
            cell = _write(worksheet, current_row, column + offset, value)
            _style(
                cell,
                JOB_AMOUNT if offset == 2 else WRAPPED_CELL if offset in {1, 4, 8} else JOB_CELL,
            )
        _height(
            worksheet,
            current_row,
            wrapped_row_height(
                report.title if report else "",
                report.location if report else "",
                characters_per_line=44,
            ),
        )


def _expenses(worksheet, model: WeeklyAccountingXlsxModel, row: int, column: int) -> None:
    for offset, label in enumerate(("NOTE", "EXPENSES", "AMOUNT"), start=9):
        cell = _write(worksheet, row, column + offset, label)
        _style(cell, EXPENSE_HEADER)
    for index, expense in enumerate(model.expenses, start=1):
        current_row = row + index
        values = (
            expense.note,
            expense.expense_type,
            expense.amount,
        )
        for offset, value in enumerate(values, start=9):
            cell = _write(worksheet, current_row, column + offset, value)
            _style(cell, EXPENSE_AMOUNT if offset == 11 else EXPENSE_CELL)
        _height(
            worksheet,
            current_row,
            wrapped_row_height(expense.note, expense.expense_type, characters_per_line=32),
        )


def _summary(worksheet, model: WeeklyAccountingXlsxModel, row: int, column: int) -> None:
    label_column = column + 10
    value_column = column + 11

    def pair(offset: int, label: str, value: object, *, header: bool = False, money: bool = False):
        label_cell = _write(worksheet, row + offset, label_column, label)
        value_cell = _write(worksheet, row + offset, value_column, value)
        _style(label_cell, SUMMARY_HEADER if header else SUMMARY_CELL)
        _style(
            value_cell,
            TOTAL_VALUE
            if header and money
            else SUMMARY_HEADER
            if header
            else SUMMARY_AMOUNT
            if money
            else SUMMARY_CELL,
        )

    pair(0, "TOTAL EXPENSES", "", header=True)
    pair(1, "Expenses", model.totals.expense_total, money=True)
    pair(2, "REVIEWS", "", header=True)
    for index, (code, count) in enumerate(model.totals.reviews, start=3):
        pair(index, REVIEW_LABELS[code], count)
    pair(6, "AMOUNTS", "", header=True)
    for index, (code, amount) in enumerate(model.totals.payments, start=7):
        pair(index, PAYMENT_LABELS[code], amount, money=True)
    pair(15, "ACTIVITY", "", header=True)
    pair(16, "Reports", model.totals.report_count)
    pair(17, "Expense entries", model.totals.expense_count)
    pair(18, "Maintenance plans", model.totals.maintenance_count)
    for index, (code, count) in enumerate(model.totals.closed_by, start=19):
        pair(index, CLOSER_LABELS[code], count)
    pair(21, "TOTAL", model.totals.gross_total, header=True, money=True)
    pair(22, "DATE", f"{model.week_start:%m/%d/%Y} - {model.week_end:%m/%d/%Y}")


def render_technician_weekly_block(
    worksheet,
    model: WeeklyAccountingXlsxModel,
    origin_row: int,
    origin_col: int,
) -> int:
    layout = block_layout(model)
    _set_widths(worksheet, origin_col)
    for day_index, relative_row in enumerate(layout.day_rows):
        _day_section(worksheet, model, day_index, origin_row + relative_row, origin_col)
    _expenses(worksheet, model, origin_row, origin_col)
    _summary(worksheet, model, origin_row + layout.summary_row, origin_col)
    for relative in range(layout.row_count):
        _height(worksheet, origin_row + relative, DATA_ROW_HEIGHT)
    return layout.row_count


def _serialize(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_individual_weekly_xlsx(model: WeeklyAccountingXlsxModel) -> bytes:
    workbook, worksheet = _new_workbook("Weekly Report", model.week_start)
    _configure(worksheet, all_tech=False)
    _title(worksheet, model, 1, 1)
    height = render_technician_weekly_block(worksheet, model, 2, 1)
    worksheet.print_area = f"A1:L{height + 1}"
    return _serialize(workbook)


def render_all_tech_weekly_xlsx(
    models: tuple[WeeklyAccountingXlsxModel, ...], week_start: date
) -> bytes:
    workbook, worksheet = _new_workbook("All Tech Weekly Report", week_start)
    _configure(worksheet, all_tech=True)
    if not models:
        _write(worksheet, 1, 1, "No technicians included for this week")
        return _serialize(workbook)
    next_row = 1
    max_row = 1
    max_column = TECHNICIAN_BLOCK_COLUMNS
    for band_start in range(0, len(models), TECHNICIANS_PER_BAND):
        band = models[band_start : band_start + TECHNICIANS_PER_BAND]
        band_height = max(block_layout(model).row_count for model in band)
        for position, model in enumerate(band):
            column = 1 + position * TECHNICIAN_COLUMN_STRIDE
            _title(worksheet, model, next_row, column)
            render_technician_weekly_block(worksheet, model, next_row + 1, column)
            worksheet.column_dimensions[
                get_column_letter(column + TECHNICIAN_BLOCK_COLUMNS)
            ].width = 3
            max_column = max(max_column, column + TECHNICIAN_BLOCK_COLUMNS - 1)
        max_row = next_row + band_height
        next_row += 1 + band_height + BAND_GAP_ROWS
    worksheet.print_area = f"A1:{get_column_letter(max_column)}{max_row}"
    return _serialize(workbook)
