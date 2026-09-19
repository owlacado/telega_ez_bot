import asyncio
import re
import time
import tracemalloc
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import event, text, update

from hub.accounting import service
from hub.accounting.domain import (
    AccountingContext,
    AccountingTotals,
    ExpenseFact,
    ReportFact,
    build_weekly,
)
from hub.accounting.service import calculate_all_weekly
from hub.accounting.xlsx import (
    XLSX_CONTENT_TYPE,
    WeeklyAccountingXlsxModel,
    block_layout,
    individual_filename,
    render_all_tech_weekly_xlsx,
    render_individual_weekly_xlsx,
)
from hub.accounting.xlsx_styles import CURRENCY_FORMAT, LEGACY_BLUE, LEGACY_YELLOW
from hub.auth.models import Manager
from hub.technicians.models import Technician
from tests.accounting_data import seed

START = date(2026, 9, 14)


def report(**values):
    return ReportFact(
        id=uuid4(),
        revision_number=values.get("revision_number", 1),
        business_date=values.get("business_date", START),
        sequence=values.get("sequence", 1),
        start_time="09:00",
        end_time="10:00",
        title=values.get("title", "Synthetic job"),
        location=values.get("location", "Synthetic location"),
        comments="Private comments are intentionally not exported",
        amount=Decimal(values.get("amount", "193.01")),
        payment_method=values.get("payment_method", "CASH"),
        closed_by=values.get("closed_by", "MYSELF"),
        google_reviews=values.get("google_reviews", 2),
        groupon_reviews=values.get("groupon_reviews", 1),
        facebook_reviews=values.get("facebook_reviews", 3),
        maintenance=values.get("maintenance", True),
        submitted_at=datetime(2026, 9, 14, 12, tzinfo=UTC),
    )


def expense(**values):
    return ExpenseFact(
        id=uuid4(),
        revision_number=values.get("revision_number", 1),
        business_date=values.get("business_date", START),
        accounting_timezone="America/Los_Angeles",
        expense_type=values.get("expense_type", "Synthetic supplies"),
        amount=Decimal(values.get("amount", "85.82")),
        note=values.get("note", "Synthetic expense note"),
        submitted_at=datetime(2026, 9, 14, 13, tzinfo=UTC),
    )


def weekly(
    *,
    technician_id: UUID | None = None,
    name: str = "Synthetic Technician",
    reports: tuple[ReportFact, ...] | None = None,
    expenses: tuple[ExpenseFact, ...] | None = None,
):
    context = AccountingContext(
        technician_id=technician_id or uuid4(),
        technician_name=name,
        accounting_timezone="America/Los_Angeles",
        setup_required=False,
        today=START,
        calculated_at=datetime(2026, 9, 18, 12, tzinfo=UTC),
    )
    return build_weekly(
        context,
        START,
        (report(),) if reports is None else reports,
        (expense(),) if expenses is None else expenses,
    )


def model(**values) -> WeeklyAccountingXlsxModel:
    return WeeklyAccountingXlsxModel.from_accounting(weekly(**values))


def workbook(content: bytes):
    return load_workbook(BytesIO(content), data_only=False, keep_links=False)


def money(cell) -> Decimal:
    return Decimal(str(cell.value)).quantize(Decimal("0.01"))


def test_individual_renderer_uses_canonical_values_and_legacy_visual_roles(record_property):
    source = model()
    started = time.perf_counter()
    content = render_individual_weekly_xlsx(source)
    elapsed = (time.perf_counter() - started) * 1000
    tracemalloc.start()
    render_individual_weekly_xlsx(source)
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    record_property("individual_ordinary_ms", str(round(elapsed, 2)))
    record_property("individual_ordinary_bytes", str(len(content)))
    record_property("individual_ordinary_peak_bytes", str(peak))
    book = workbook(content)
    sheet = book.active
    assert book.sheetnames == ["Weekly Report"]
    assert sheet["A1"].value == "Synthetic Technician"
    assert sheet["A2"].value == "MONDAY  09/14/2026"
    assert sheet["A3"].value == "#"
    assert sheet["B4"].value == "Synthetic job\nSynthetic location"
    assert money(sheet["C4"]) == Decimal("193.01")
    assert sheet["C4"].number_format == CURRENCY_FORMAT
    assert sheet["J2"].value == "NOTE"
    assert sheet["K3"].value == "Synthetic supplies"
    assert money(sheet["L3"]) == Decimal("85.82")
    assert sheet["K39"].value == "REVIEWS"
    assert [sheet[f"K{row}"].value for row in (40, 41, 42)] == [
        "Google",
        "Groupon",
        "Facebook",
    ]
    assert sheet["K43"].value == "AMOUNTS"
    assert [sheet[f"K{row}"].value for row in range(44, 52)] == [
        "Cash",
        "Zelle",
        "Check",
        "Credit Card / Cash App",
        "Venmo",
        "SUPER",
        "Estimate",
        "Cancel",
    ]
    assert sheet["K58"].value == "TOTAL"
    assert money(sheet["L58"]) == Decimal("193.01")
    assert sheet["A2"].fill.fgColor.rgb == f"FF{LEGACY_BLUE}"
    assert sheet["A3"].fill.fgColor.rgb == f"FF{LEGACY_YELLOW}"
    assert sheet["C4"].fill.patternType is None
    assert sheet["A121"].value is None
    assert "Private comments" not in [cell.value for row in sheet.iter_rows() for cell in row]


@pytest.mark.parametrize("scenario", ["empty", "expense_only", "outcomes_only"])
def test_empty_expense_only_and_outcome_only_weeks(scenario):
    reports = ()
    expenses = ()
    if scenario == "expense_only":
        expenses = (expense(amount="20.00"),)
    if scenario == "outcomes_only":
        reports = (
            report(sequence=1, amount="0.00", payment_method="ESTIMATE"),
            report(sequence=2, amount="0.00", payment_method="CANCEL"),
        )
    sheet = workbook(
        render_individual_weekly_xlsx(model(reports=reports, expenses=expenses))
    ).active
    assert sheet["A2"].value.startswith("MONDAY")
    assert sheet["A104"].value.startswith("SUNDAY")
    assert money(sheet["L58"]) == Decimal("0.00")
    assert money(sheet["L38"]) == (
        Decimal("20.00") if scenario == "expense_only" else Decimal("0.00")
    )
    if scenario == "expense_only":
        assert sheet["K3"].value == "Synthetic supplies"
        assert sheet["A4"].value is None
    if scenario == "outcomes_only":
        assert [sheet[f"D{row}"].value for row in (4, 5)] == ["Estimate", "Cancel"]
        assert [money(sheet[f"C{row}"]) for row in (4, 5)] == [
            Decimal("0.00"),
            Decimal("0.00"),
        ]


def test_renderer_uses_supplied_totals_without_recalculating_rows():
    canonical = weekly()
    supplied = AccountingTotals(
        gross_total=Decimal("777.77"),
        expense_total=Decimal("66.66"),
        report_count=41,
        expense_count=42,
        maintenance_count=43,
        payments={
            "CASH": Decimal("1.01"),
            "ZELLE": Decimal("2.02"),
            "CHECK": Decimal("3.03"),
            "CREDIT_CARD": Decimal("4.04"),
            "VENMO": Decimal("5.05"),
            "SUPER": Decimal("6.06"),
            "ESTIMATE": Decimal("0.00"),
            "CANCEL": Decimal("0.00"),
        },
        reviews={"GOOGLE": 51, "GROUPON": 52, "FACEBOOK": 53},
        closed_by={"MYSELF": 54, "CALL_CENTER": 55},
    )
    content = render_individual_weekly_xlsx(
        WeeklyAccountingXlsxModel.from_accounting(canonical.model_copy(update={"totals": supplied}))
    )
    sheet = workbook(content).active
    assert money(sheet["L38"]) == Decimal("66.66")
    assert [sheet[f"L{row}"].value for row in (40, 41, 42)] == [51, 52, 53]
    assert [money(sheet[f"L{row}"]) for row in range(44, 52)] == [
        Decimal("1.01"),
        Decimal("2.02"),
        Decimal("3.03"),
        Decimal("4.04"),
        Decimal("5.05"),
        Decimal("6.06"),
        Decimal("0.00"),
        Decimal("0.00"),
    ]
    assert [sheet[f"L{row}"].value for row in range(53, 58)] == [41, 42, 43, 54, 55]
    assert money(sheet["L58"]) == Decimal("777.77")
    assert money(sheet["C4"]) == Decimal("193.01")


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_user_text_remains_literal_and_control_characters_are_removed(prefix):
    hostile = f'{prefix}HYPERLINK("https://evil.invalid","click")'
    item = report(title=hostile, location=hostile + "\nUnicode Привет 😀\x01")
    cost = expense(expense_type=hostile, note=hostile + "\x02")
    content = render_individual_weekly_xlsx(
        model(name=hostile + "\x03", reports=(item,), expenses=(cost,))
    )
    sheet = workbook(content).active
    for address in ("A1", "B4", "J3", "K3"):
        cell = sheet[address]
        assert cell.data_type == "s"
        assert cell.data_type != "f"
        assert "HYPERLINK" in cell.value
        assert not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", cell.value)
        assert cell.hyperlink is None
    assert "Привет 😀" in sheet["B4"].value


def test_workbook_is_self_contained_without_formulas_macros_or_external_links():
    content = render_individual_weekly_xlsx(model())
    book = workbook(content)
    assert not book._external_links
    assert all(
        cell.data_type != "f" and cell.hyperlink is None
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
        assert not any("externalLinks/" in name for name in names)
        assert not any("vbaProject" in name or name.endswith(".bin") for name in names)
        relationships = "".join(
            archive.read(name).decode("utf-8", "replace")
            for name in names
            if name.endswith(".rels")
        )
        assert 'TargetMode="External"' not in relationships


@pytest.mark.parametrize(
    "value",
    [
        Decimal("0.01"),
        Decimal("1.10"),
        Decimal("193.01"),
        Decimal("1793.16"),
        Decimal("9999999999.99"),
        Decimal("99999999999900.00"),
    ],
)
def test_exact_money_roundtrip_and_large_width(value):
    source = weekly()
    totals = source.totals.model_copy(update={"gross_total": value})
    rendered = render_individual_weekly_xlsx(
        WeeklyAccountingXlsxModel.from_accounting(source.model_copy(update={"totals": totals}))
    )
    sheet = workbook(rendered).active
    assert money(sheet["L58"]) == value
    assert sheet["L58"].number_format == CURRENCY_FORMAT
    assert sheet.column_dimensions["L"].width >= 24


def test_long_text_and_one_hundred_reports_and_expenses_expand_without_truncation(
    record_property,
):
    reports = tuple(
        report(
            sequence=index + 1,
            title=f"Synthetic job {index}",
            location="Long synthetic address " * 20,
            amount="0.01",
        )
        for index in range(100)
    )
    expenses = tuple(
        expense(
            expense_type=f"Synthetic type {index}",
            note="Long synthetic expense note " * 20,
            amount="0.01",
        )
        for index in range(100)
    )
    large = model(reports=reports, expenses=expenses)
    assert block_layout(large).row_count > 119
    started = time.perf_counter()
    content = render_individual_weekly_xlsx(large)
    elapsed = (time.perf_counter() - started) * 1000
    tracemalloc.start()
    render_individual_weekly_xlsx(large)
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    record_property("individual_100_100_ms", str(round(elapsed, 2)))
    record_property("individual_100_100_bytes", str(len(content)))
    record_property("individual_100_100_peak_bytes", str(peak))
    sheet = workbook(content).active
    assert sheet["A103"].value == 100
    assert {sheet.cell(row, 11).value for row in range(3, 103)} == {
        f"Synthetic type {index}" for index in range(100)
    }
    assert sheet["B4"].alignment.wrap_text is True
    assert sheet["J3"].alignment.wrap_text is True
    assert sheet.row_dimensions[4].height > 18
    assert sheet.max_row >= block_layout(large).row_count + 1


def test_individual_and_all_tech_share_values_styles_and_widths():
    current = model()
    individual = workbook(render_individual_weekly_xlsx(current)).active
    consolidated = workbook(render_all_tech_weekly_xlsx((current,), START)).active
    for address in ("A1", "A2", "A3", "B4", "C4", "J3", "K39", "L58"):
        assert individual[address].value == consolidated[address].value
        assert individual[address]._style == consolidated[address]._style
    for column in "ABCDEFGHIJKL":
        assert (
            individual.column_dimensions[column].width
            == consolidated.column_dimensions[column].width
        )


@pytest.mark.parametrize("count", [1, 10, 11, 20, 21, 30])
def test_all_tech_has_no_ten_technician_cap(count, record_property):
    models = tuple(
        model(technician_id=UUID(int=index + 1), name=f"Tech {index + 1:02d}")
        for index in range(count)
    )
    started = time.perf_counter()
    content = render_all_tech_weekly_xlsx(models, START)
    elapsed = (time.perf_counter() - started) * 1000
    record_property(f"all_tech_{count}_ms", str(round(elapsed, 2)))
    record_property(f"all_tech_{count}_bytes", str(len(content)))
    if count == 30:
        tracemalloc.start()
        render_all_tech_weekly_xlsx(models, START)
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        record_property("all_tech_30_peak_bytes", str(peak))
    sheet = workbook(content).active
    assert sheet["A1"].value == "Tech 01"
    if count >= 11:
        assert sheet["A124"].value == "Tech 11"
    if count >= 21:
        assert sheet["A247"].value == "Tech 21"
    assert (
        sum(
            1
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("Tech ")
        )
        == count
    )


def test_dynamic_tall_block_cannot_overlap_next_band():
    tall = model(
        technician_id=UUID(int=1),
        name="Tall Tech",
        reports=tuple(report(sequence=i + 1, amount="0.01") for i in range(100)),
    )
    ordinary = tuple(
        model(technician_id=UUID(int=i + 2), name=f"Ordinary {i + 2}") for i in range(10)
    )
    sheet = workbook(render_all_tech_weekly_xlsx((tall, *ordinary), START)).active
    expected_second_band = 1 + 1 + block_layout(tall).row_count + 3
    assert sheet.cell(expected_second_band, 1).value == "Ordinary 11"
    assert sheet.cell(expected_second_band - 1, 1).value is None


def test_safe_filename_is_ascii_bounded_and_header_safe():
    unsafe = model(name='../ "\r\n=HYPERLINK("x") Привет')
    filename = individual_filename(unsafe)
    assert filename.endswith("_2026-09-14_2026-09-20.xlsx")
    assert len(filename) < 140
    assert re.fullmatch(r"[A-Za-z0-9_.-]+", filename)
    assert ".." not in filename and "\r" not in filename and "\n" not in filename


async def test_individual_endpoint_auth_validation_and_reopen(client, app):
    technician_id = await seed(app.state.session_factory, START)
    path = f"/api/technicians/{technician_id}/accounting/weekly.xlsx"
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://127.0.0.1:3000"
    ) as unauthenticated:
        assert (await unauthenticated.get(path, params={"week_start": START})).status_code == 401
        assert (
            await unauthenticated.get(
                path,
                params={"week_start": START},
                headers={"Authorization": "Bearer synthetic-capability"},
            )
        ).status_code == 401
    for query in (
        "week_start=2026-09-15",
        "week_start=2026-02-30",
        "week_start=",
        "week_start=America%2FLos_Angeles",
        "week_start=2026-09-14&week_start=2026-09-21",
    ):
        response = await client.get(f"{path}?{query}")
        assert response.status_code == 422
        assert response.headers["cache-control"] == "no-store"
    response = await client.get(path, params={"week_start": START})
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_CONTENT_TYPE
    assert response.headers["cache-control"] == "no-store"
    assert re.fullmatch(
        r'attachment; filename="[A-Za-z0-9_.-]+\.xlsx"',
        response.headers["content-disposition"],
    )
    assert workbook(response.content).active["L58"].value == 1793.16
    assert (
        await client.get(
            f"/api/technicians/{uuid4()}/accounting/weekly.xlsx",
            params={"week_start": START},
        )
    ).status_code == 404


async def test_endpoint_preserves_latest_only_and_identical_expenses(client, app):
    technician_id = await seed(
        app.state.session_factory,
        START,
        reports=[
            {
                "day": 0,
                "amount": "25.00",
                "previous_amount": "999.99",
                "revision_number": 2,
                "payment_method": "ZELLE",
            }
        ],
        expenses=[
            {
                "day": 0,
                "amount": "20.00",
                "previous_amount": "888.88",
                "revision_number": 2,
            },
            {"day": 0, "amount": "20.00"},
        ],
    )
    response = await client.get(
        f"/api/technicians/{technician_id}/accounting/weekly.xlsx",
        params={"week_start": START},
    )
    sheet = workbook(response.content).active
    assert money(sheet["C4"]) == Decimal("25.00")
    assert money(sheet["L45"]) == Decimal("25.00")
    assert [money(sheet[f"L{row}"]) for row in (3, 4)] == [
        Decimal("20.00"),
        Decimal("20.00"),
    ]
    assert money(sheet["L38"]) == Decimal("40.00")
    assert sheet["L53"].value == 1
    assert sheet["L54"].value == 2
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert 999.99 not in values
    assert 888.88 not in values


async def test_inactive_manager_cannot_download(client, app, credentials):
    technician_id = await seed(app.state.session_factory, START, reports=[], expenses=[])
    async with app.state.session_factory() as db, db.begin():
        await db.execute(
            update(Manager).where(Manager.id == credentials["id"]).values(is_active=False)
        )
    response = await client.get(
        f"/api/technicians/{technician_id}/accounting/weekly.xlsx",
        params={"week_start": START},
    )
    assert response.status_code == 401
    assert response.headers["cache-control"] == "no-store"


async def test_generation_failure_is_sanitized_and_business_content_is_not_logged(
    client, app, monkeypatch, caplog
):
    technician_id = await seed(app.state.session_factory, START)
    path = f"/api/technicians/{technician_id}/accounting/weekly.xlsx"
    caplog.set_level("INFO")
    success = await client.get(path, params={"week_start": START})
    assert success.status_code == 200
    assert "Synthetic location" not in caplog.text
    assert "synthetic expense" not in caplog.text

    def broken(model):
        raise RuntimeError("PRIVATE_PATH PRIVATE_EXPENSE_NOTE")

    monkeypatch.setattr("hub.accounting.xlsx_router.render_individual_weekly_xlsx", broken)
    failure = await client.get(path, params={"week_start": START})
    assert failure.status_code == 500
    assert failure.headers["cache-control"] == "no-store"
    assert failure.json() == {
        "error": {
            "code": "internal_error",
            "message": "Unable to complete request. Retry safely.",
        }
    }
    assert "PRIVATE_PATH" not in caplog.text
    assert "PRIVATE_EXPENSE_NOTE" not in caplog.text


async def test_download_does_not_create_business_audit_events(client, app):
    technician_id = await seed(app.state.session_factory, START)
    async with app.state.session_factory() as db:
        before = len((await db.scalars(text("SELECT id FROM audit_events"))).all())
    response = await client.get(
        f"/api/technicians/{technician_id}/accounting/weekly.xlsx",
        params={"week_start": START},
    )
    assert response.status_code == 200
    async with app.state.session_factory() as db:
        after = len((await db.scalars(text("SELECT id FROM audit_events"))).all())
    assert after == before


async def test_all_tech_inclusion_order_snapshot_and_query_budget(app):
    ids = [uuid4() for _ in range(5)]
    async with app.state.session_factory() as db, db.begin():
        db.add_all(
            [
                Technician(
                    id=ids[0],
                    first_name="Zulu",
                    last_name="Active Empty",
                    status="ACTIVE",
                    accounting_timezone="America/Los_Angeles",
                ),
                Technician(
                    id=ids[1],
                    first_name="Alpha",
                    last_name="Inactive History",
                    status="INACTIVE",
                    accounting_timezone="America/New_York",
                ),
                Technician(
                    id=ids[2],
                    first_name="Hidden",
                    last_name="Inactive Empty",
                    status="INACTIVE",
                    accounting_timezone="America/Los_Angeles",
                ),
                Technician(
                    id=ids[3],
                    first_name="Bravo",
                    last_name="Active History",
                    status="ACTIVE",
                    accounting_timezone=None,
                ),
                Technician(
                    id=ids[4],
                    first_name="Future",
                    last_name="Active Empty",
                    status="ACTIVE",
                    accounting_timezone="America/Los_Angeles",
                ),
            ]
        )
        await db.execute(
            update(Technician)
            .where(Technician.id == ids[4])
            .values(created_at=datetime(2026, 10, 1, tzinfo=UTC))
        )
    await seed(
        app.state.session_factory,
        START,
        technician_id=ids[1],
        reports=[{"day": 0, "amount": "1.00", "payment_method": "CASH"}],
        expenses=[],
    )
    await seed(
        app.state.session_factory,
        START,
        technician_id=ids[3],
        reports=[],
        expenses=[{"day": 1, "amount": "2.00"}],
    )
    statements = []

    def track(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(app.state.engine.sync_engine, "before_cursor_execute", track)
    try:
        async with app.state.session_factory() as db:
            result = await calculate_all_weekly(db, START)
            export_statements = len(statements)
            isolation = await db.scalar(text("SHOW transaction_isolation"))
            read_only = await db.scalar(text("SHOW transaction_read_only"))
    finally:
        event.remove(app.state.engine.sync_engine, "before_cursor_execute", track)
    assert export_statements == 4
    assert isolation == "repeatable read"
    assert read_only == "on"
    assert [item.technician_name for item in result] == [
        "Alpha Inactive History",
        "Bravo Active History",
        "Zulu Active Empty",
    ]
    assert [item.totals.gross_total for item in result] == [Decimal("1.00"), 0, 0]
    assert [item.totals.expense_total for item in result] == [0, Decimal("2.00"), 0]
    assert result[1].setup_required is True


async def test_all_tech_endpoint_reopens_and_preserves_inactive_history(client, app):
    technician_id = await seed(app.state.session_factory, START)
    async with app.state.session_factory() as db, db.begin():
        await db.execute(
            update(Technician).where(Technician.id == technician_id).values(status="INACTIVE")
        )
    response = await client.get("/api/accounting/weekly/all.xlsx", params={"week_start": START})
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_CONTENT_TYPE
    sheet = workbook(response.content).active
    assert sheet["A1"].value == "Accounting Synthetic"
    assert money(sheet["L58"]) == Decimal("1793.16")


async def test_all_tech_database_query_count_is_four_for_thirty_technicians(app):
    async with app.state.session_factory() as db, db.begin():
        db.add_all(
            [
                Technician(
                    first_name="Scale",
                    last_name=f"Tech {index:02d}",
                    accounting_timezone="America/Los_Angeles",
                )
                for index in range(30)
            ]
        )
    statements = []

    def track(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(app.state.engine.sync_engine, "before_cursor_execute", track)
    try:
        async with app.state.session_factory() as db:
            result = await calculate_all_weekly(db, START)
    finally:
        event.remove(app.state.engine.sync_engine, "before_cursor_execute", track)
    assert len(result) == 30
    assert len(statements) == 4


async def test_all_tech_uses_one_coherent_snapshot_during_concurrent_submission(app, monkeypatch):
    technician_id = await seed(app.state.session_factory, START)
    original = service.all_facts

    async def interleaved(db, start, end):
        assert await db.scalar(text("SHOW transaction_isolation")) == "repeatable read"
        assert await db.scalar(text("SHOW transaction_read_only")) == "on"
        await asyncio.wait_for(
            seed(
                app.state.session_factory,
                START,
                technician_id=technician_id,
                reports=[{"day": 0, "amount": "1.00", "payment_method": "CASH"}],
                expenses=[{"day": 0, "amount": "2.00"}],
            ),
            5,
        )
        return await original(db, start, end)

    monkeypatch.setattr(service, "all_facts", interleaved)
    async with app.state.session_factory() as db:
        result = await calculate_all_weekly(db, START)
    assert len(result) == 1
    assert result[0].totals.gross_total == Decimal("1793.16")
    assert result[0].totals.expense_total == Decimal("150.82")
    monkeypatch.setattr(service, "all_facts", original)
    async with app.state.session_factory() as db:
        after = await calculate_all_weekly(db, START)
    assert after[0].totals.gross_total == Decimal("1794.16")
    assert after[0].totals.expense_total == Decimal("152.82")
