"""Independent adversarial audit coverage for Stage 8 weekly XLSX exports."""

import asyncio
import inspect
import re
import time
import tracemalloc
import zipfile
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import event, text, update

from hub.accounting import service
from hub.accounting.domain import (
    AccountingContext,
    ExpenseFact,
    ReportFact,
    build_weekly,
)
from hub.accounting.service import calculate_all_weekly
from hub.accounting.xlsx import (
    WeeklyAccountingXlsxModel,
    block_layout,
    individual_filename,
    render_all_tech_weekly_xlsx,
    render_individual_weekly_xlsx,
    spreadsheet_text,
)
from hub.accounting.xlsx_styles import CURRENCY_FORMAT
from hub.auth.models import Manager, ManagerSession
from hub.technicians.models import Technician
from tests.accounting_data import seed

START = date(2026, 9, 14)
PAYMENTS = (
    ("CASH", "1.01"),
    ("ZELLE", "2.02"),
    ("CHECK", "3.03"),
    ("CREDIT_CARD", "4.04"),
    ("VENMO", "5.05"),
    ("SUPER", "6.06"),
    ("ESTIMATE", "0.00"),
    ("CANCEL", "0.00"),
)


def _report(index: int, payment: str = "CASH", amount: str = "1.00", **values):
    return ReportFact(
        id=values.get("id", uuid4()),
        revision_number=values.get("revision_number", 2),
        business_date=values.get("business_date", START),
        sequence=index,
        start_time="08:00",
        end_time="09:00",
        title=values.get("title", f"Audit job {index}"),
        location=values.get("location", f"Audit location {index}"),
        comments="Not exported",
        amount=Decimal(amount),
        payment_method=payment,
        closed_by=values.get("closed_by", "MYSELF" if index % 2 else "CALL_CENTER"),
        google_reviews=values.get("google_reviews", 20),
        groupon_reviews=values.get("groupon_reviews", 21),
        facebook_reviews=values.get("facebook_reviews", 22),
        maintenance=values.get("maintenance", index % 2 == 1),
        submitted_at=datetime(2026, 9, 14, 12, index % 60, tzinfo=UTC),
    )


def _expense(index: int, amount: str = "1.00", **values):
    return ExpenseFact(
        id=values.get("id", uuid4()),
        revision_number=values.get("revision_number", 2),
        business_date=values.get("business_date", START),
        accounting_timezone=values.get("timezone", "America/Los_Angeles"),
        expense_type=values.get("expense_type", f"Audit expense {index}"),
        amount=Decimal(amount),
        note=values.get("note", f"Audit note {index}"),
        submitted_at=datetime(2026, 9, 14, 13, index % 60, tzinfo=UTC),
    )


def _weekly(
    *,
    identifier: UUID | None = None,
    name: str = "Audit Technician",
    timezone: str = "America/Los_Angeles",
    reports: tuple[ReportFact, ...] = (),
    expenses: tuple[ExpenseFact, ...] = (),
):
    context = AccountingContext(
        technician_id=identifier or uuid4(),
        technician_name=name,
        accounting_timezone=timezone,
        setup_required=False,
        today=START,
        calculated_at=datetime(2026, 9, 18, 12, tzinfo=UTC),
    )
    return build_weekly(context, START, reports, expenses)


def _model(**values):
    return WeeklyAccountingXlsxModel.from_accounting(_weekly(**values))


def _book(content: bytes):
    return load_workbook(BytesIO(content), data_only=False, keep_links=False)


def _money(value) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


def _all_values(sheet):
    return [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]


def test_audit_renderer_is_a_presentation_only_boundary():
    source = inspect.getsource(__import__("hub.accounting.xlsx", fromlist=["xlsx"]))
    assert "sqlalchemy" not in source
    assert "WorkReport" not in source
    assert "ExpenseRevision" not in source
    assert "select(" not in source
    assert "=SUM" not in source.upper()
    assert "gross_total -" not in source
    assert "float(" not in source


def test_audit_individual_and_all_tech_match_every_canonical_total_and_style():
    reports = tuple(
        _report(index, payment, amount) for index, (payment, amount) in enumerate(PAYMENTS, start=1)
    )
    expenses = (_expense(1, "7.07"), _expense(2, "8.08"))
    canonical = _weekly(reports=reports, expenses=expenses)
    model = WeeklyAccountingXlsxModel.from_accounting(canonical)
    individual = _book(render_individual_weekly_xlsx(model)).active
    all_tech = _book(render_all_tech_weekly_xlsx((model,), START)).active

    assert canonical.totals.gross_total == Decimal("21.21")
    assert canonical.totals.expense_total == Decimal("15.15")
    assert [_money(individual[f"L{row}"].value) for row in range(44, 52)] == [
        amount for amount in canonical.totals.payments.values()
    ]
    assert [individual[f"L{row}"].value for row in (40, 41, 42)] == list(
        canonical.totals.reviews.values()
    )
    assert [individual[f"L{row}"].value for row in range(53, 58)] == [
        canonical.totals.report_count,
        canonical.totals.expense_count,
        canonical.totals.maintenance_count,
        *canonical.totals.closed_by.values(),
    ]
    assert _money(individual["L38"].value) == canonical.totals.expense_total
    assert _money(individual["L58"].value) == canonical.totals.gross_total
    for address in ("A1", "L1", "A2", "A3", "B4", "C4", "J3", "K39", "L58"):
        assert individual[address].value == all_tech[address].value
        assert individual[address]._style == all_tech[address]._style
    for column in "ABCDEFGHIJKL":
        assert (
            individual.column_dimensions[column].width == all_tech.column_dimensions[column].width
        )
    for row in (1, 2, 3, 4, 39, 58):
        assert individual.row_dimensions[row].height == all_tech.row_dimensions[row].height


@pytest.mark.parametrize(
    "payload",
    [
        '=HYPERLINK("https://evil.invalid","CLICK")',
        '=WEBSERVICE("https://evil.invalid")',
        "=1+1",
        "+1+1",
        "-2+3",
        "@SUM(A1:A2)",
        "=CMD|' /C calc'!A0",
        "@SUM(1+1)*cmd|' /C calc'!A0",
        '=IMPORTXML("https://evil.invalid","//x")',
        "http://evil.invalid/path",
        "https://evil.invalid/path",
        "file:///C:/private.txt",
        "javascript:alert(1)",
        "mailto:a@example.invalid",
        r"\\server\share\file.xlsx",
        r"C:\private\file.xlsx",
    ],
)
def test_audit_formula_dde_uri_and_hyperlink_payloads_remain_plain_text(payload):
    model = _model(
        name=payload,
        reports=(_report(1, title=payload, location=payload),),
        expenses=(_expense(1, expense_type=payload, note=payload),),
    )
    content = render_individual_weekly_xlsx(model)
    book = _book(content)
    for address in ("A1", "B4", "J3", "K3"):
        cell = book.active[address]
        assert cell.data_type == "s"
        assert payload in cell.value
        assert cell.hyperlink is None
    assert not book._external_links


def test_audit_xml_controls_noncharacters_and_unicode_round_trip():
    prohibited = (
        "".join(chr(value) for value in (*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20), 0x7F))
        + "\ud800\udfff\ufffe\uffff"
    )
    legitimate = (
        "Cyrillic Привет · emoji 😀 · café e\u0301 · RTL مرحبا · zero\u200bwidth\tline\nnext"
    )
    cleaned = spreadsheet_text(f"{prohibited}{legitimate}{prohibited}")
    assert all(character not in cleaned for character in prohibited)
    assert cleaned == legitimate
    book = _book(
        render_individual_weekly_xlsx(
            _model(
                name=f"{prohibited}{legitimate}",
                reports=(_report(1, title=legitimate, location=legitimate),),
                expenses=(_expense(1, expense_type=legitimate, note=legitimate),),
            )
        )
    )
    assert book.active["A1"].value == legitimate
    assert legitimate in book.active["B4"].value
    assert book.active["J3"].value == legitimate


@pytest.mark.parametrize(
    "amount",
    [
        "0.00",
        "0.01",
        "0.10",
        "1.10",
        "193.01",
        "1793.16",
        "9999999999.99",
        "99999999999900.00",
        "-193.01",
    ],
)
def test_audit_money_round_trip_format_width_and_negative_passthrough(amount):
    expected = Decimal(amount)
    presentation = _model(reports=(_report(1),))
    first_day = presentation.days[0]
    presentation = replace(
        presentation,
        days=(
            replace(
                first_day,
                reports=(replace(first_day.reports[0], amount=expected),),
            ),
            *presentation.days[1:],
        ),
        totals=replace(presentation.totals, gross_total=expected),
    )
    sheet = _book(render_individual_weekly_xlsx(presentation)).active
    assert _money(sheet["C4"].value) == expected
    assert _money(sheet["L58"].value) == expected
    assert sheet["C4"].number_format == CURRENCY_FORMAT
    assert sheet["L58"].number_format == CURRENCY_FORMAT
    assert sheet.column_dimensions["L"].width >= 24


def test_audit_empty_expense_only_and_outcome_only_semantics():
    cases = (
        _weekly(),
        _weekly(expenses=(_expense(1, "12.34"),)),
        _weekly(
            reports=(
                _report(1, "ESTIMATE", "0.00"),
                _report(2, "CANCEL", "0.00"),
            )
        ),
    )
    for canonical in cases:
        sheet = _book(
            render_individual_weekly_xlsx(WeeklyAccountingXlsxModel.from_accounting(canonical))
        ).active
        assert [sheet.cell(2 + index * 17, 1).value.split()[0] for index in range(7)] == [
            "MONDAY",
            "TUESDAY",
            "WEDNESDAY",
            "THURSDAY",
            "FRIDAY",
            "SATURDAY",
            "SUNDAY",
        ]
        assert _money(sheet["L58"].value) == canonical.totals.gross_total
        assert _money(sheet["L38"].value) == canonical.totals.expense_total
        assert not any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)


def test_audit_identical_expenses_and_latest_values_are_not_deduplicated():
    same = dict(
        amount="25.00",
        expense_type="Identical",
        note="Exactly identical",
        business_date=START,
    )
    canonical = _weekly(
        reports=(_report(1, "ZELLE", "150.00", revision_number=2),),
        expenses=(_expense(1, **same), _expense(2, **same)),
    )
    sheet = _book(
        render_individual_weekly_xlsx(WeeklyAccountingXlsxModel.from_accounting(canonical))
    ).active
    assert [sheet[f"K{row}"].value for row in (3, 4)] == ["Identical", "Identical"]
    assert [_money(sheet[f"L{row}"].value) for row in (3, 4)] == [
        Decimal("25.00"),
        Decimal("25.00"),
    ]
    assert _money(sheet["L45"].value) == Decimal("150.00")


@pytest.mark.parametrize("count", [31, 50])
def test_audit_additional_fleet_bands_have_no_cap_duplicates_or_overlap(count, record_property):
    models = tuple(
        _model(identifier=UUID(int=index + 1), name=f"Audit Tech {index + 1:03d}")
        for index in range(count)
    )
    started = time.perf_counter()
    content = render_all_tech_weekly_xlsx(models, START)
    elapsed = (time.perf_counter() - started) * 1000
    record_property(f"audit_all_tech_{count}_ms", str(round(elapsed, 2)))
    record_property(f"audit_all_tech_{count}_bytes", str(len(content)))
    sheet = _book(content).active
    values = _all_values(sheet)
    names = [f"Audit Tech {index + 1:03d}" for index in range(count)]
    assert all(values.count(name) == 1 for name in names)
    assert sheet["A370"].value == "Audit Tech 031"
    if count == 50:
        assert sheet["A493"].value == "Audit Tech 041"
        assert sheet.cell(493, 118).value == "Audit Tech 050"
    assert sheet.max_column == 129
    assert sheet.max_row <= 612


def test_audit_optional_one_hundred_technician_probe(record_property):
    models = tuple(
        _model(identifier=UUID(int=index + 1), name=f"Hundred {index + 1:03d}")
        for index in range(100)
    )
    tracemalloc.start()
    started = time.perf_counter()
    content = render_all_tech_weekly_xlsx(models, START)
    elapsed = (time.perf_counter() - started) * 1000
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    sheet = _book(content).active
    assert _all_values(sheet).count("Hundred 100") == 1
    assert sheet["A1108"].value == "Hundred 091"
    record_property("audit_all_tech_100_ms", str(round(elapsed, 2)))
    record_property("audit_all_tech_100_bytes", str(len(content)))
    record_property("audit_all_tech_100_peak_bytes", str(peak))


def test_audit_performance_and_memory_recheck(record_property):
    ordinary = _model(reports=(_report(1, amount="193.01"),), expenses=(_expense(1),))
    large = _model(
        reports=tuple(_report(index, amount="0.01") for index in range(1, 101)),
        expenses=tuple(_expense(index, "0.01") for index in range(1, 101)),
    )

    def measure(name, render):
        started = time.perf_counter()
        content = render()
        record_property(f"{name}_ms", str(round((time.perf_counter() - started) * 1000, 2)))
        record_property(f"{name}_bytes", str(len(content)))

    measure("audit_individual_ordinary", lambda: render_individual_weekly_xlsx(ordinary))
    measure("audit_individual_100_100", lambda: render_individual_weekly_xlsx(large))
    for count in (10, 20, 30, 50):
        models = tuple(
            _model(identifier=UUID(int=index + 1), name=f"Profile {index + 1:03d}")
            for index in range(count)
        )
        measure(
            f"audit_all_tech_{count}",
            lambda models=models: render_all_tech_weekly_xlsx(models, START),
        )
        if count in {30, 50}:
            tracemalloc.start()
            render_all_tech_weekly_xlsx(models, START)
            _, peak = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            record_property(f"audit_all_tech_{count}_peak_bytes", str(peak))


def test_audit_five_hundred_rows_and_maximum_text_keep_geometry():
    reports = tuple(
        _report(
            index,
            amount="0.01",
            title=f"Job {index}",
            location=("L" * 2000 if index == 1 else f"Location {index}"),
        )
        for index in range(1, 501)
    )
    expenses = tuple(
        _expense(
            index,
            "0.01",
            expense_type=f"Type {index}",
            note=("N" * 4000 if index == 1 else f"Note {index}"),
        )
        for index in range(1, 501)
    )
    model = _model(reports=reports, expenses=expenses)
    assert block_layout(model).row_count == 604
    sheet = _book(render_individual_weekly_xlsx(model)).active
    assert sheet["A503"].value == 500
    assert {sheet[f"K{row}"].value for row in range(3, 503)} == {
        f"Type {index}" for index in range(1, 501)
    }
    assert sheet["K503"].value == "TOTAL EXPENSES"
    assert sheet["A504"].value.startswith("TUESDAY")
    assert len(sheet["B4"].value.splitlines()[1]) == 2000
    long_expense_row = next(row for row in range(3, 503) if sheet[f"K{row}"].value == "Type 1")
    assert len(sheet[f"J{long_expense_row}"].value) == 4000
    assert 18 < sheet.row_dimensions[4].height <= 180
    assert 18 < sheet.row_dimensions[long_expense_row].height <= 180


def test_audit_zip_geometry_hidden_state_metadata_and_relationships():
    models = tuple(
        _model(identifier=UUID(int=index + 1), name=f"Geometry {index + 1:02d}")
        for index in range(50)
    )
    content = render_all_tech_weekly_xlsx(models, START)
    book = _book(content)
    sheet = book.active
    assert book.sheetnames == ["All Tech Weekly Report"]
    assert sheet.sheet_state == "visible"
    assert not sheet.merged_cells.ranges
    assert not any(dimension.hidden for dimension in sheet.row_dimensions.values())
    assert not any(dimension.hidden for dimension in sheet.column_dimensions.values())
    assert book.properties.creator == "Technician Hub"
    metadata = " ".join(
        str(value)
        for value in (
            book.properties.creator,
            book.properties.title,
            book.properties.subject,
            book.properties.lastModifiedBy,
        )
        if value
    )
    assert not re.search(r"rasha|HVAC_TECHNICIAN_HUB|postgresql|127\.0\.0\.1|token", metadata, re.I)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
        assert all("\\" not in name and not re.match(r"^[A-Za-z]:", name) for name in names)
        assert not any(
            token.lower() in name.lower()
            for name in names
            for token in ("externalLinks", "vbaProject", "oleObject", ".bin")
        )
        xml = "\n".join(
            archive.read(name).decode("utf-8", "replace")
            for name in names
            if name.endswith((".xml", ".rels"))
        )
        assert 'TargetMode="External"' not in xml
        assert not re.search(r"\bDDE\b|vbaProject|oleObject", xml, re.I)


@pytest.mark.parametrize(
    "name",
    [
        "../name",
        r"..\name",
        'quote"name',
        "single'name",
        "line\r\nbreak;name",
        "emoji😀name",
        "Я" * 200,
        "A" * 1000,
    ],
)
def test_audit_filename_and_header_component_is_ascii_bounded(name):
    filename = individual_filename(_model(name=name))
    assert re.fullmatch(r"[A-Za-z0-9_.-]+", filename)
    assert len(filename) <= 116
    assert ".." not in filename
    assert "\r" not in filename and "\n" not in filename
    assert "/" not in filename and "\\" not in filename


@pytest.mark.parametrize("count", [1, 10, 30, 50])
async def test_audit_all_tech_query_budget_stays_four(app, count):
    async with app.state.session_factory() as db, db.begin():
        db.add_all(
            [
                Technician(
                    id=UUID(int=index + 1),
                    first_name="Query",
                    last_name=f"Tech {index:03d}",
                    accounting_timezone="America/Los_Angeles",
                )
                for index in range(count)
            ]
        )
        await db.flush()
        await db.execute(update(Technician).values(created_at=datetime(2026, 9, 14, tzinfo=UTC)))
    statements = []

    def track(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(app.state.engine.sync_engine, "before_cursor_execute", track)
    try:
        async with app.state.session_factory() as db:
            result = await calculate_all_weekly(db, START)
    finally:
        event.remove(app.state.engine.sync_engine, "before_cursor_execute", track)
    assert len(result) == count
    assert len(statements) == 4


async def test_audit_mixed_timezones_inclusion_order_and_current_name_semantics(app):
    ids = [UUID(int=index) for index in range(1, 9)]
    rows = [
        ("Alice", "Same", "ACTIVE", "America/New_York"),
        ("alice", "Same", "ACTIVE", "America/Chicago"),
        ("Álice", "Same", "ACTIVE", "America/Denver"),
        ("Same", "Name", "ACTIVE", "America/Los_Angeles"),
        ("Same", "Name", "ACTIVE", "America/Phoenix"),
        ("Inactive", "History", "INACTIVE", "America/New_York"),
        ("Inactive", "Empty", "INACTIVE", "America/Chicago"),
        ("Future", "Active", "ACTIVE", "America/Denver"),
    ]
    async with app.state.session_factory() as db, db.begin():
        db.add_all(
            [
                Technician(
                    id=identifier,
                    first_name=first,
                    last_name=last,
                    status=status,
                    accounting_timezone=zone,
                )
                for identifier, (first, last, status, zone) in zip(ids, rows, strict=True)
            ]
        )
        await db.flush()
        await db.execute(
            update(Technician)
            .where(Technician.id != ids[7])
            .values(created_at=datetime(2026, 9, 14, tzinfo=UTC))
        )
    await seed(
        app.state.session_factory,
        START,
        technician_id=ids[5],
        reports=[{"day": 0, "amount": "5.00", "payment_method": "CASH"}],
        expenses=[],
    )
    async with app.state.session_factory() as db, db.begin():
        await db.execute(
            update(Technician)
            .where(Technician.id == ids[7])
            .values(created_at=datetime(2026, 9, 21, tzinfo=UTC))
        )
        await db.execute(
            update(Technician).where(Technician.id == ids[5]).values(first_name="Renamed")
        )
    async with app.state.session_factory() as db:
        result = await calculate_all_weekly(db, START)
    names = [item.technician_name for item in result]
    assert names == sorted(names, key=lambda value: value.casefold())
    assert names[:2] == ["Alice Same", "alice Same"]
    assert "Álice Same" in names
    assert "Renamed History" in names
    assert "Inactive Empty" not in names
    assert "Future Active" not in names
    assert {item.accounting_timezone for item in result} >= {
        "America/New_York",
        "America/Chicago",
        "America/Denver",
        "America/Los_Angeles",
        "America/Phoenix",
    }
    assert all(
        [day.business_date for day in item.days]
        == [
            date(2026, 9, 14),
            date(2026, 9, 15),
            date(2026, 9, 16),
            date(2026, 9, 17),
            date(2026, 9, 18),
            date(2026, 9, 19),
            date(2026, 9, 20),
        ]
        for item in result
    )


async def test_audit_fleet_snapshot_is_coherent_and_does_not_block_two_writers(app, monkeypatch):
    first = await seed(
        app.state.session_factory,
        START,
        reports=[{"day": 0, "amount": "10.00", "payment_method": "CASH"}],
        expenses=[],
    )
    second = await seed(
        app.state.session_factory,
        START,
        reports=[{"day": 0, "amount": "20.00", "payment_method": "CASH"}],
        expenses=[],
    )
    original = service.all_facts

    async def interleaved(db, start, end):
        assert await db.scalar(text("SHOW transaction_isolation")) == "repeatable read"
        assert await db.scalar(text("SHOW transaction_read_only")) == "on"
        await asyncio.wait_for(
            asyncio.gather(
                seed(
                    app.state.session_factory,
                    START,
                    technician_id=first,
                    reports=[{"day": 0, "amount": "1.00", "payment_method": "ZELLE"}],
                    expenses=[{"day": 0, "amount": "3.00"}],
                ),
                seed(
                    app.state.session_factory,
                    START,
                    technician_id=second,
                    reports=[{"day": 0, "amount": "2.00", "payment_method": "ZELLE"}],
                    expenses=[{"day": 0, "amount": "4.00"}],
                ),
            ),
            timeout=5,
        )
        return await original(db, start, end)

    monkeypatch.setattr(service, "all_facts", interleaved)
    async with app.state.session_factory() as db:
        during = await calculate_all_weekly(db, START)
    assert sorted(item.totals.gross_total for item in during) == [
        Decimal("10.00"),
        Decimal("20.00"),
    ]
    monkeypatch.setattr(service, "all_facts", original)
    async with app.state.session_factory() as db:
        after = await calculate_all_weekly(db, START)
    assert sorted(item.totals.gross_total for item in after) == [
        Decimal("11.00"),
        Decimal("22.00"),
    ]
    assert sorted(item.totals.expense_total for item in after) == [
        Decimal("3.00"),
        Decimal("4.00"),
    ]


async def test_audit_renderer_executes_zero_database_queries_after_model_creation(app):
    technician_id = await seed(app.state.session_factory, START)
    async with app.state.session_factory() as db:
        accounting = await service.calculate(db, technician_id, "weekly", START)
    model = WeeklyAccountingXlsxModel.from_accounting(accounting)
    statements = []

    def track(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(app.state.engine.sync_engine, "before_cursor_execute", track)
    try:
        assert render_individual_weekly_xlsx(model)
        assert render_all_tech_weekly_xlsx((model,), START)
    finally:
        event.remove(app.state.engine.sync_engine, "before_cursor_execute", track)
    assert statements == []


async def test_audit_both_routes_authorization_mime_and_inactive_history(client, app, credentials):
    technician_id = await seed(app.state.session_factory, START)
    individual = f"/api/technicians/{technician_id}/accounting/weekly.xlsx"
    all_tech = "/api/accounting/weekly/all.xlsx"
    params = {"week_start": START.isoformat()}
    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://127.0.0.1:3000",
        headers={"Origin": "http://127.0.0.1:3000", "X-Hub-Request": "1"},
    ) as unauthenticated:
        for path in (individual, all_tech):
            denied = await unauthenticated.get(path, params=params)
            assert denied.status_code == 401
            capability = await unauthenticated.get(
                path,
                params=params,
                headers={"Authorization": "Bearer audit-technician-capability"},
            )
            assert capability.status_code == 401
            allowed = await client.get(path, params=params)
            assert allowed.status_code == 200
            assert (
                allowed.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert allowed.headers["cache-control"] == "no-store"
            assert allowed.headers["content-disposition"].startswith('attachment; filename="')
            assert _book(allowed.content).active

        async with app.state.session_factory() as db, db.begin():
            await db.execute(
                update(Technician).where(Technician.id == technician_id).values(status="INACTIVE")
            )
        historical = await client.get(individual, params=params)
        assert historical.status_code == 200
        assert historical.headers["cache-control"] == "no-store"

        async with app.state.session_factory() as db, db.begin():
            await db.execute(
                update(ManagerSession).values(expires_at=datetime.now(UTC) - timedelta(seconds=1))
            )
        for path in (individual, all_tech):
            assert (await client.get(path, params=params)).status_code == 401

        login = await unauthenticated.post(
            "/api/auth/login",
            json={
                "username": credentials["username"],
                "password": credentials["password"],
            },
        )
        assert login.status_code == 200
        unauthenticated.headers["X-CSRF-Token"] = login.json()["csrf_token"]
        async with app.state.session_factory() as db, db.begin():
            await db.execute(
                update(Manager).where(Manager.id == credentials["id"]).values(is_active=False)
            )
        for path in (individual, all_tech):
            assert (await unauthenticated.get(path, params=params)).status_code == 401
